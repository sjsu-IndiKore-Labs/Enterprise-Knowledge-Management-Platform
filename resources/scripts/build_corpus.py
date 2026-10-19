"""Render resources/corpus/source/*.md into the PDF/DOCX files listed in manifest.csv,
plus the edge-case files in resources/edge-cases/.

Usage (from repo root):
    python3 -m venv .venv && .venv/bin/pip install -r resources/scripts/requirements.txt
    .venv/bin/python resources/scripts/build_corpus.py            # formats from manifest
    .venv/bin/python resources/scripts/build_corpus.py --all-formats   # also every doc as both PDF and DOCX

Markdown subset understood: '# ' title, '## ' heading, '- ' bullets, '1. ' numbered items,
'| a | b |' tables, '<!-- pagebreak -->' explicit page break (page numbers in eval/golden_qa.jsonl
depend on these breaks), the 'Document ID: ...' metadata line, plain paragraphs.
"""
import csv
import os
import re
import sys
from pathlib import Path

from docx import Document as DocxDocument
from docx.enum.text import WD_BREAK
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (ListFlowable, ListItem, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table,
                                TableStyle)

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "corpus" / "source"
OUT = ROOT / "corpus" / "files"
EDGE = ROOT / "edge-cases"


def parse(md_text):
    """Return list of blocks: (kind, payload)."""
    blocks, para, table, items, list_kind = [], [], [], [], None

    def flush():
        nonlocal para, table, items, list_kind
        if para:
            blocks.append(("p", " ".join(para)))
            para = []
        if table:
            blocks.append(("table", table))
            table = []
        if items:
            blocks.append((list_kind, items))
            items, list_kind = [], None

    for raw in md_text.splitlines():
        line = raw.rstrip()
        if line.strip() == "<!-- pagebreak -->":
            flush(); blocks.append(("pagebreak", None)); continue
        if not line.strip():
            flush(); continue
        if line.startswith("# "):
            flush(); blocks.append(("title", line[2:].strip())); continue
        if line.startswith("## "):
            flush(); blocks.append(("h2", line[3:].strip())); continue
        if line.startswith("Document ID:") or line.startswith("Change note:"):
            flush(); blocks.append(("meta", line.strip())); continue
        if line.startswith("|"):
            if para or items:
                flush()
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                continue
            table.append(cells); continue
        m_b = re.match(r"^- (.*)", line)
        m_n = re.match(r"^\d+\. (.*)", line)
        if m_b or m_n:
            kind = "ul" if m_b else "ol"
            if table or para or (items and list_kind != kind):
                flush()
            list_kind = kind
            items.append((m_b or m_n).group(1)); continue
        if table or items:
            flush()
        para.append(line.strip())
    flush()
    return blocks


def esc(s):
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def build_pdf(blocks, path, footer_title):
    styles = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=styles["BodyText"], fontSize=10, leading=13.5, spaceAfter=6)
    meta = ParagraphStyle("meta", parent=body, fontSize=8, textColor=colors.grey)
    cell = ParagraphStyle("cell", parent=body, fontSize=9, leading=11, spaceAfter=0)
    story = []
    for kind, val in blocks:
        if kind == "title":
            story.append(Paragraph(esc(val), styles["Title"]))
        elif kind == "h2":
            story.append(Paragraph(esc(val), styles["Heading2"]))
        elif kind == "meta":
            story.append(Paragraph(esc(val), meta)); story.append(Spacer(1, 6))
        elif kind == "p":
            story.append(Paragraph(esc(val), body))
        elif kind in ("ul", "ol"):
            story.append(ListFlowable([ListItem(Paragraph(esc(i), body)) for i in val],
                                      bulletType="bullet" if kind == "ul" else "1", leftIndent=14))
        elif kind == "table":
            data = [[Paragraph(esc(c), cell) for c in row] for row in val]
            t = Table(data, hAlign="LEFT", repeatRows=1)
            t.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E5E5E5")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]))
            story.append(t); story.append(Spacer(1, 8))
        elif kind == "pagebreak":
            story.append(PageBreak())

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.grey)
        canvas.drawString(inch, 0.5 * inch, f"Helios Systems Inc. — Internal — {footer_title}")
        canvas.drawRightString(letter[0] - inch, 0.5 * inch, f"Page {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(str(path), pagesize=letter, topMargin=0.8 * inch, bottomMargin=0.9 * inch,
                            title=footer_title, author="Helios Systems Inc. (fictional)")
    doc.build(story, onFirstPage=footer, onLaterPages=footer)


def build_docx(blocks, path):
    d = DocxDocument()
    for kind, val in blocks:
        if kind == "title":
            d.add_heading(val, level=0)
        elif kind == "h2":
            d.add_heading(val, level=1)
        elif kind in ("meta", "p"):
            d.add_paragraph(val)
        elif kind in ("ul", "ol"):
            for i in val:
                d.add_paragraph(i, style="List Bullet" if kind == "ul" else "List Number")
        elif kind == "table":
            t = d.add_table(rows=len(val), cols=len(val[0]))
            t.style = "Table Grid"
            for r, row in enumerate(val):
                for c, txt in enumerate(row):
                    t.cell(r, c).text = txt
        elif kind == "pagebreak":
            d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    d.core_properties.author = "Helios Systems Inc. (fictional)"
    d.save(str(path))


def pdf_page_count(path):
    return len(re.findall(rb"/Type\s*/Page[^s]", path.read_bytes()))


def build_edge_cases():
    from reportlab.pdfgen import canvas
    from reportlab.lib.pdfencrypt import StandardEncryption
    EDGE.mkdir(exist_ok=True)
    # 1. image-only PDF: shapes but no text layer (simulates a scanned page)
    c = canvas.Canvas(str(EDGE / "EDGE-001_scanned_no_text_layer.pdf"), pagesize=letter)
    for i in range(40):
        c.setFillGray(0.2 + (i % 5) * 0.1)
        c.rect(inch, 10 * inch - i * 0.22 * inch, (3 + (i * 7) % 4) * inch, 0.12 * inch, fill=1, stroke=0)
    c.save()
    # 2. corrupt PDF (valid header, garbage body)
    (EDGE / "EDGE-002_corrupt.pdf").write_bytes(b"%PDF-1.7\n" + os.urandom(2048))
    # 3. unsupported type
    (EDGE / "EDGE-003_unsupported_format.txt").write_text("Plain text files are not accepted; upload must return 415.\n")
    # 4. large multi-page PDF (load / progress testing)
    blocks = [("title", "Helios Platform Engineering Manual (Synthetic, 60 pages)"),
              ("meta", "Document ID: EDGE-004 | Department: ENGINEERING | Category: MANUAL | Synthetic load-test document")]
    for p in range(1, 61):
        blocks.append(("h2", f"Chapter {p}: Service module {p:02d}"))
        for k in range(5):
            blocks.append(("p", f"Module {p:02d} section {k + 1}. The service module {p:02d} exposes a health endpoint, "
                                f"publishes metrics every {10 + p % 7} seconds, and retries failed calls up to {2 + p % 4} "
                                f"times with exponential backoff starting at {100 * (1 + p % 3)} milliseconds. Owners rotate "
                                f"quarterly and the runbook identifier for this module is RB-{1000 + p}."))
        if p < 60:
            blocks.append(("pagebreak", None))
    build_pdf(blocks, EDGE / "EDGE-004_large_60_pages.pdf", "Platform Engineering Manual")
    # 5. empty DOCX
    DocxDocument().save(str(EDGE / "EDGE-005_empty.docx"))
    # 6. password-protected PDF
    enc = StandardEncryption("helios-secret", canPrint=0)
    c = canvas.Canvas(str(EDGE / "EDGE-006_password_protected.pdf"), pagesize=letter, encrypt=enc)
    c.drawString(inch, 10 * inch, "This PDF is encrypted. Ingestion should fail with CORRUPT_FILE or UNSUPPORTED_FORMAT.")
    c.save()


def main():
    all_formats = "--all-formats" in sys.argv
    OUT.mkdir(parents=True, exist_ok=True)
    with open(ROOT / "corpus" / "manifest.csv", newline="") as f:
        rows = list(csv.DictReader(f))
    for row in rows:
        blocks = parse((SRC / row["sourceMarkdown"]).read_text())
        stem = Path(row["file"]).stem
        formats = {"pdf", "docx"} if all_formats else {row["format"]}
        if "pdf" in formats:
            p = OUT / f"{stem}.pdf"
            build_pdf(blocks, p, row["title"])
            print(f"{p.name:70s} {pdf_page_count(p)} pages")
        if "docx" in formats:
            p = OUT / f"{stem}.docx"
            build_docx(blocks, p)
            print(f"{p.name:70s} docx")
    build_edge_cases()
    print("edge cases ->", EDGE)


if __name__ == "__main__":
    main()
